import os as _os
ROOT = _os.environ.get('PAC_ROOT') or _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
import csv, json, re, sys, collections, types
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

D0 = f'{ROOT}/pac-data/'
OUT = sys.argv[1]
tr = json.load(open(f'{ROOT}/keldaancommunity/pokemonautochess/app/public/dist/client/locales/en/translation.json'))
allp = list(csv.DictReader(open(D0 + 'pokemons-data.csv')))
core = list(csv.DictReader(open(D0 + 'core-value-tiers.csv')))

# ---------- scale (provisional) ----------
def hp_c(x): return max(3, round(x / 10))
def atk_c(x): return max(1, round(x / 5))
def df_c(x): return max(0, round(x / 7))
def spd_c(x): return max(1, round(x / 7))
def pp_c(x): return max(2, round(x / 35))
def D(x): return max(1, int(x / 10 + 0.5))          # flat PAC damage / shield -> game
def PA(p, base): return max(1, int(p / 100 * base / 6 + 0.5))  # percent of a PAC stat
def S(x): return min(3, max(1, int(x / 3 + 0.5)))   # seconds -> rounds
def AT(x): return max(1, int(x / 5 + 0.5))
def SP(x): return max(1, int(x / 7 + 0.5))

import re as _re, sys as _s2
_s2.path.insert(0, f'{ROOT}/pac-sim')
from duel import APM as _APM
def apmark(ab, txt):
    # AP symbol next to every number AP boosts ("AP/2" = +1 per 2 AP, rounded down)
    mode = _APM.get(ab, 'first')
    if mode == 'none': return txt
    tag = '\u2726' if mode == 'first' else '\u2726/2 '
    if ab == 'DRUM_BEATING':
        txt = _re.sub(r'gain (\d+) SHIELD', lambda m: 'gain ' + tag + m.group(1) + ' SHIELD', txt, 1)
        return _re.sub(r'deal (\d+) SPECIAL', lambda m: 'deal ' + tag + m.group(1) + ' SPECIAL', txt, 1)
    for pat in (r'(\d+) (?:SPECIAL|TRUE|PHYSICAL)', r'hits of (\d+)', r'(?i:gain) (\d+) SHIELD', r'Heal (\d+)'):
        m = _re.search(pat, txt)
        if m:
            return txt[:m.start(1)] + tag + txt[m.start(1):]
    return txt

# ---------- charge powers ----------
# template(v, c) ; v = star-resolved numbers in order of the PAC arrays ; c = PAC stats
SPL = " SPLASH."
P = {
'ACCELEROCK': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL. Gain +1 SPEED and -1 DEF (max 3 times)."),
'ACID_ARMOR': ('Tank', lambda v, c: f"Gain +{AT(v[0])} DEF. The next 2 attackers that hit you lose 1 DEF each."),
'AGILITY': ('Speedster', lambda v, c: f"Gain +{SP(v[0])} SPEED for the rest of the duel."),
'AIR_SLASH': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and FLINCH the target (1 turn)."),
'AQUA_STEP': ('Speedster', lambda v, c: f"Gain +{SP(v[1])} SPEED, then deal {D(v[0])} SPECIAL."),
'AQUA_TAIL': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and gain {D(v[1])} SHIELD."),
'BITE': ('Drain', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL, heal half the damage dealt, and FLINCH the target (1 turn)."),
'BLAST_BURN': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'BLAZE_KICK': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and BURN the target. +30% damage (round up) if it was already burning."),
'BLOOD_MOON': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL and WOUND the target (2 turns)."),
'BUG_BUZZ': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, doubled if the target is PARALYZED."),
'BULLDOZE': ('Control', lambda v, c: f"Deal {D(v[0]*1.4)} SPECIAL and give the target -1 SPEED." + SPL),
'COLUMN_CRUSH': ('Striker', lambda v, c: f"Gain {D(v[0])} SHIELD, then deal SPECIAL equal to your total SHIELD."),
'CRABHAMMER': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL. KO the target if it is at 30% HP or less."),
'CRUNCH': ('Drain', lambda v, c: f"Deal {D(v[0])} SPECIAL. If it KOs the target, heal half the target's max HP."),
'DARKEST_LARIAT': ('MultiHit', lambda v, c: f"Hit {max(2, round(c.spd/25))} times for {PA(v[0], c.atk)} SPECIAL each (more hits with high SPEED) and FLINCH the target."),
'DARK_HARVEST': ('Drain', lambda v, c: f"Deal {D(v[0]*3)} SPECIAL and heal {max(1, round(D(v[0]*3)*0.3))}. You are FLINCHED on your next turn."),
'DOUBLE_SHOCK': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL. You are PARALYZED for 1 turn."),
'DRAGON_BREATH': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'DRAGON_TAIL': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and gain +{AT(v[1])} DEF and SP.DEF."),
'DRUM_BEATING': ('Support', lambda v, c: f"Cycles each cast: (1) gain {D(v[0]*.8)} SHIELD; (2) deal {D(v[1]*.8)} SPECIAL; (3) your next Pokemon enters with +{SP(v[2])} SPEED."),
'ENTANGLING_THREAD': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and PARALYZE the target (2 turns)." + SPL),
'FAIRY_WIND': ('Support', lambda v, c: f"Your next Pokemon enters with +{1 + (1 if v[0] >= 10 else 0)} charge."),
'FIRESTARTER': ('Speedster', lambda v, c: f"Gain +{SP(v[0])} SPEED, then deal {D(v[1]*.8)} SPECIAL."),
'FLAMETHROWER': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and BURN the target." + SPL),
'FLOWER_TRICK': ('Striker', lambda v, c: f"Plant a bouquet: after your next turn the target takes {D(v[0])} SPECIAL (+1 per crit it suffers meanwhile)."),
'FURY_SWIPES': ('MultiHit', lambda v, c: f"Make {max(1, round(v[0]/5))} extra basic attack(s) (ATK minus DEF each). If one KOs the target, gain +1 charge."),
'FUTURE_SIGHT': ('Striker', lambda v, c: f"After your 2nd turn from now, the opposing Pokemon takes {D(v[1]*3)} SPECIAL."),
'GEAR_GRIND': ('MultiHit', lambda v, c: f"Fire two gears: 2 x {PA(v[0], c.spd)} SPECIAL (scales with SPEED)."),
'GIGATON_HAMMER': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL. You are FATIGUED (no charge gain) for 2 turns."),
'GLAIVE_RUSH': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL. You and the target both get ARMOR BREAK (2 turns)."),
'GROWL': ('Control', lambda v, c: f"FLINCH the target (1 turn) and give it -{AT(v[0])} ATK (2 turns)."),
'GUILLOTINE': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL. If it KOs the target, gain +1 charge."),
'HEADBUTT': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and FLINCH the target (1 turn). Doubled if the target has SHIELD."),
'HEAVY_SLAM': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, +1 for every 3 max HP you have above the target's." + SPL),
'HEX': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, doubled if the target has a negative status."),
'HORN_ATTACK': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL and ARMOR BREAK the target (1 turn)."),
'HORN_DRILL': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL, +50% (round up) if your ATK is higher than the target's."),
'HYDRO_PUMP': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'ICE_BALL': ('Tank', lambda v, c: f"Gain +2 SP.DEF, then deal {D(v[0]) + PA(v[1], c.sdef)} SPECIAL (scales with your SP.DEF)."),
'ICICLE_CRASH': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'ICICLE_MISSILE': ('Control', lambda v, c: f"Launch {int(v[0])} icicle(s) for {D(v[1]*1.5)} SPECIAL each." + (" The last one FREEZES the target (1 turn)." if v[0] >= 3 else "")),
'ICY_WIND': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and give the target -{SP(v[1])} SPEED." + SPL),
'KING_SHIELD': ('Tank', lambda v, c: f"PROTECT (ignore all damage until your next turn) and gain {D(v[0]*2)} SHIELD."),
'KOWTOW_CLEAVE': ('Striker', lambda v, c: f"Deal {int(PA(150, c.atk)*1.5+0.5)} SPECIAL (always crit), +{max(1, round(PA(150, c.atk)*v[0]/100))} TRUE per fallen Pokemon of yours."),
'LEAF_BLADE': ('Striker', lambda v, c: f"Deal {int(PA(v[0], c.atk)*1.5+0.5)} TRUE (always crit, ignores all blocking)."),
'LEECH_LIFE': ('Drain', lambda v, c: f"Deal {D(v[0])} SPECIAL and heal the same amount."),
'LICK': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and PARALYZE + CONFUSE the target (1 turn)."),
'MAGICAL_LEAF': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and ARMOR BREAK the target (1 turn)."),
'MAGIC_POWDER': ('Tank', lambda v, c: f"Gain {D(v[0])} SHIELD and FLINCH the target ({S(v[1])} turn(s))."),
'MAGNET_BOMB': ('Control', lambda v, c: f"Deal {D(v[0]*1.5)} SPECIAL and give the target -1 SPEED (LOCKED: it cannot FLY AWAY or dodge until its next turn)." + f" SPLASH {D(v[1])}."),
'MANTIS_BLADES': ('MultiHit', lambda v, c: f"Three hits of {D(v[0])}: one PHYSICAL, one SPECIAL, one TRUE."),
'METEOR_MASH': ('MultiHit', lambda v, c: f"Hit 3 times (4 with a PSYCHIC synergy) for {PA(v[0], c.atk)} SPECIAL each, then gain +1 ATK."),
'MYSTICAL_FIRE': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and give the target -1 AP."),
'NIGHTMARE': ('Control', lambda v, c: f"FATIGUE the target ({S(v[0])} turn(s)). If it has a negative status, also deal {D(v[1])} SPECIAL."),
'NUZZLE': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and PARALYZE the target ({S(v[1])} turn)."),
'PECK': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL."),
'PETAL_DANCE': ('MultiHit', lambda v, c: f"Release {max(2, round(v[0]/2))} petals for {D(v[1])} SPECIAL each."),
'PLAY_ROUGH': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and CHARM the target (-2 ATK, 2 turns)."),
'PSYCHIC': ('Control', lambda v, c: f"Deal {D(v[0]*1.5)} SPECIAL; the target loses 1 charge." + SPL),
'PSYCHO_CUT': ('MultiHit', lambda v, c: f"Three blades for {int(D(v[0])*1.5+0.5)} SPECIAL each (crit)."),
'RAPID_SPIN': ('Tank', lambda v, c: f"Deal {D(v[0])} SPECIAL and gain +{max(1, round(PA(v[1], c.atk)))} DEF and SP.DEF."),
'REFLECT': ('Tank', lambda v, c: f"For {S(v[0])} turn(s), all PHYSICAL damage you take is blocked and {int(v[1])}% of it is reflected as SPECIAL."),
'RETALIATE': ('Striker', lambda v, c: f"Deal {PA(v[0], c.atk)} SPECIAL, plus one more equal hit for each fallen Pokemon of yours."),
'ROCK_ARTILLERY': ('MultiHit', lambda v, c: f"Throw {max(2, round(v[0]/5))} rocks for {D(v[1]*1.5)} SPECIAL each."),
'ROCK_SLIDE': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, doubled if the target is FLYING."),
'SALT_CURE': ('Support', lambda v, c: f"Gain {D(v[0])} SHIELD and cure your statuses. If the target is WATER, STEEL or GHOST, BURN it."),
'SHADOW_BALL': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL and give the target -1 SP.DEF (stacks)."),
'SHOCKWAVE': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'SILVER_WIND': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL. Gain +1 ATK and +1 SPEED."),
'SING': ('Control', lambda v, c: f"Put the target to SLEEP ({S(v[1])} turn(s)): it skips its turns."),
'SLASH': ('Striker', lambda v, c: f"Deal {int(D(v[0])*1.5+0.5)} SPECIAL (always crit)."),
'SNIPE_SHOT': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, ignoring SP.DEF."),
'SOAK': ('Support', lambda v, c: f"Deal {D(v[0])} SPECIAL. Your next Pokemon enters with +1 charge."),
'SOFT_BOILED': ('Support', lambda v, c: f"Cure your statuses and gain {D(v[0])} SHIELD. Your next Pokemon enters with half that SHIELD (round up)."),
'SPIKY_SHIELD': ('Tank', lambda v, c: f"For {S(v[0])} turn(s), whoever hits you with a PHYSICAL attack is WOUNDED and takes {PA(v[1], c.defn)} SPECIAL."),
'STEAMROLLER': ('Striker', lambda v, c: f"Deal {PA(v[0]*.75, c.spd)} SPECIAL (scales with SPEED)."),
'STORED_POWER': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL, +1 for each stat boost you have (ATK, SPEED, DEF, SP.DEF, AP)."),
'STRING_SHOT': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and PARALYZE the target (2 turns)."),
'TELEPORT': ('Striker', lambda v, c: f"Your next attack this duel deals +{D(v[0])} extra SPECIAL."),
'TERRAIN_PULSE': ('Support', lambda v, c: f"Heal {max(1, round(v[0]/100*c.hpc))}. With GRASS gain +1 DEF; with ELECTRIC gain +1 SPEED; with PSYCHIC gain +1 AP."),
'THRASH': ('Striker', lambda v, c: f"Gain +{max(1, round(v[0]/100*c.atkc))} ATK for the duel, then CONFUSE yourself (your next attack is halved)."),
'THUNDER_SHOCK': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL."),
'TICKLE': ('Control', lambda v, c: f"Give the target -1 ATK and -1 DEF (2 turns)."),
'TORCH_SONG': ('MultiHit', lambda v, c: f"Four flames for {PA(50, c.atk)} SPECIAL each; the last BURNs the target. Gain +{int(v[2])} AP."),
'TRANSE': ('Support', lambda v, c: f"Heal 50% of max HP (min 1) and return to normal form."),
'TRI_ATTACK': ('Control', lambda v, c: f"Deal {D(v[0]*.8)} SPECIAL and choose: BURN (2 turns), PARALYZE (1 turn) or FREEZE (1 turn)."),
'TROP_KICK': ('Control', lambda v, c: f"Deal {D(v[0])} SPECIAL and give the target -{AT(v[1])} ATK (2 turns)."),
'TWISTER': ('Striker', lambda v, c: f"Deal {D(v[0])} SPECIAL." + SPL),
'UPROAR': ('Striker', lambda v, c: f"Deal {D(v[0]*3)} SPECIAL. You cannot be put to SLEEP until your next turn."),
'VOLT_SWITCH': ('Speedster', lambda v, c: f"Deal {D(v[0])} SPECIAL, ignoring SP.DEF."),
'WAVE_SPLASH': ('Tank', lambda v, c: f"Gain SHIELD equal to {int(v[0])}% of max HP ({max(1, round(v[0]/100*c.hpc))}), then deal that much SPECIAL."),
'WHEEL_OF_FIRE': ('MultiHit', lambda v, c: f"Deal {D(v[0])} SPECIAL twice."),
'WHIRLPOOL': ('MultiHit', lambda v, c: f"Four whirlpools: 4 x {PA(v[0], c.atk)} SPECIAL."),
'WISH': ('Tank', lambda v, c: f"Gain {D(v[0])} SHIELD and PROTECT (ignore all damage) for 1 turn."),
}

def parse_arrays(desc):
    out = []
    for m in re.findall(r'\[([^\]]*)\]', desc):
        nums = []
        for t in m.split(','):
            t = t.strip()
            try: nums.append(float(t))
            except ValueError: pass
        if nums: out.append(nums)
    return out

def resolve(ab, stars):
    arrs = parse_arrays(tr['ability_description'].get(ab, ''))
    return [a[min(stars - 1, len(a) - 1)] for a in arrs]

# ---------- synergies (draft): thresholds = PAC thresholds x0.75 ----------
SYN = {
'NORMAL': ("SHIELD on entry", "5 SHIELD", "7 SHIELD", "9 SHIELD", "11 SHIELD, +1 ATK, +1 AP"),
'GRASS': ("Heal every turn; overheal grows max HP (max +3)", "Heal 1 at end of each of its turns", "Heal 2 at end of each of its turns", "Heal 2 per turn; overheal gives +1 max HP", "Heal 3 per turn; overheal gives +1 max HP (max +5)"),
'FIRE': ("Attacks BURN on a schedule; ATK grows", "Every 3rd attack BURNS the target (1 TRUE per turn, 3 turns)", "Every 2nd attack burns; +1 ATK after each charge power", "Every 2nd attack burns for 2 TRUE per turn; +1 ATK at the end of every one of its turns (max +4)", "Every attack burns for 2 TRUE per turn, lasting 4 turns; ATK growth max +6"),
'WATER': ("Charge faster", "Enter with +2 charge", "Enter with +2 charge; +1 charge each time you cast", "Attacks give +2 charge instead of +1"),
'ELECTRIC': ("Basic attacks trigger extra hits", "Extra hit of 1 SPECIAL after each attack", "Extra hit of 2 SPECIAL after each attack", "Two extra hits of 2 SPECIAL after each attack, and the target loses 1 charge"),
'FIGHTING': ("Block a flat amount; throw attackers away", "+1 DEF", "+1 DEF", "+2 DEF; every 3rd hit you take destroys the attacker's SHIELD and deals 1 TRUE", "+2 DEF; every 2nd hit you take destroys the attacker's SHIELD and deals 1 TRUE"),
'PSYCHIC': ("Gain AP (+1 damage to the first hit of each charge power per AP)", "+2 AP", "+3 AP", "+4 AP; your charge powers ignore SP.DEF"),
'DARK': ("Critical hits (x1.5)", "Every 3rd attack is a crit", "Every 2nd attack is a crit", "Every 2nd attack is a crit and crits are x2"),
'STEEL': ("Part of your attack damage is TRUE", "1 of your attack damage is TRUE", "2 of your attack damage is TRUE", "3 of your attack damage is TRUE", "4 of your attack damage is TRUE"),
'GROUND': ("Dig in: grow tougher the longer you stay", "+1 DEF at end of every 2nd turn (max 2)", "+1 DEF at end of every 2nd turn (max 2)", "+1 DEF and +1 ATK at end of every turn (max 4)", "+1 DEF and +1 ATK at the end of every turn (max 6)"),
'POISON': ("Attacks POISON the target", "Attacks POISON (1 stack: 1 true dmg per turn)", "Each attack adds a stack (max 3)", "Max 5 stacks"),
'DRAGON': ("Copy synergies; gain more with more Dragons", "Each Dragon's second synergy counts twice (not the third)", "Also: each Dragon enters with SHIELD equal to the number of Dragons in your lineup", "Also: each Dragon enters with +SPEED and +AP equal to half the number of Dragons (rounded down, at least 1)"),
'FIELD': ("Grow stronger as Field allies fall", "For each Field unit of yours KO'd this battle, each Field card entering gains +2 max HP (and enters at full HP)", "+3 max HP and +1 SPEED per KO'd Field unit (SPEED max +3)", "+4 max HP and +1 SPEED per KO'd Field unit (SPEED max +3)"),
'MONSTER': ("Flinch and grow on KO", "Your first attack after entry FLINCHES the target", "Also: on a KO gain +1 ATK and heal 1", "Also: +1 max HP on a KO", "On a KO gain +2 ATK, heal 1 and +1 max HP"),
'HUMAN': ("Lifesteal (heal for a share of damage dealt)", "Heal 20% of the damage you deal with attacks and charge powers (fractions carry over)", "Heal 30%", "Heal 45%"),
'AQUATIC': ("Resist status; tidal wave", "Ignore the first negative status you receive", "Also: on its 3rd turn on the field, tidal wave: cure yourself, deal 2 TRUE to the foe", "Tidal wave deals 4 and heals you 2", "Tidal wave deals 6 and heals you 4"),
'BUG': ("Swarm tokens block attacks", "1 SWARM token (blocks the next attack completely)", "2 SWARM tokens", "3 SWARM tokens; a spent token deals 1 TRUE back", "4 SWARM tokens; a spent token deals 1 TRUE back"),
'FLYING': ("Fly away at low HP", "Once, at or below 50% HP: PROTECT for 1 turn and +1 SPEED", "Same, and gain +1 SP.DEF", "Can trigger twice", "Can trigger three times"),
'FLORA': ("When KO'd, pass a BATON of growth", "BATON: your next Pokemon enters with +3 max HP", "+4 max HP", "+5 max HP and +1 ATK", "+6 max HP, +1 ATK and +1 SPEED"),
'ROCK': ("Gain DEF and resist crits", "+2 DEF", "+3 DEF; crits against you are normal hits", "+4 DEF; crits against you are normal hits"),
'GHOST': ("Evasion; curse on attack", "Every 4th attack against you misses", "Also: your first hit on each foe CURSES it (-1 DEF)", "Every 3rd attack misses; the curse also gives -1 SP.DEF and -1 ATK", "Every 2nd attack misses and deals 1 TRUE back; the curse also FATIGUES the foe for 2 turns"),
'FAIRY': ("Wands: every Nth attack casts a status wand", "Every 2nd attack CONFUSES the target", "Also CHARMS it (-2 ATK for 2 turns)", "Also slows it (-1 SPEED for 2 turns)", "Every attack casts the wand"),
'ICE': ("Freeze chance; SP.DEF", "+1 SP.DEF", "+2 SP.DEF; every 6th attack FREEZES (1 turn)", "+2 SP.DEF; every 5th attack FREEZES (1 turn)", "+3 SP.DEF; every 4th attack FREEZES (1 turn)"),
'FOSSIL': ("Awaken at low HP", "Once, at or below 30% HP: +2 SHIELD, +1 ATK", "+4 SHIELD, +1 ATK", "+6 SHIELD, +2 ATK"),
'SOUND': ("Cries empower you and the next ally", "+1 ATK on your first cast", "+1 ATK on each of your first 2 casts", "+1 ATK on each of your first 3 casts, +1 SPEED once, and every cast passes a BATON of +1 charge to your next Pokemon"),
'ARTIFICIAL': ("Stronger when holding items", "+1 ATK per held item", "+1 ATK and +1 SHIELD per held item", "+1 ATK, +1 SHIELD and +1 AP per held item"),
'BABY': ("Comfort after losses (health and gold)", "+2 max HP, plus +1 per current losing streak (max +3)", "+4 max HP, plus streak bonus", "+4 max HP, streak bonus, and when you lose this duel gain 1 gold"),
'LIGHT': ("Spotlight: your opening Pokemon (any type) is empowered", "Opening Pokemon: +1 ATK, +1 AP", "Also +2 charge", "Also +1 DEF, +1 SP.DEF and PROTECT until its first turn"),
'WILD': ("Fast and wounding", "Attacks WOUND (1 turn); +1 SPEED", "+1 SPEED, +1 ATK", "WOUND lasts 2 turns; +2 SPEED, +2 ATK", "Once at or below 30% HP: +2 ATK, +2 SPEED and 3 SHIELD"),
'AMORPHOUS': ("Whole lineup gains from variety of synergies", "Every card in your lineup gets +1 max HP per active synergy you have", "+1 max HP per active synergy and +1 SPEED per 3 active synergies", "+2 max HP per active synergy and +1 SPEED per 2 active synergies"),
'GOURMET': ("A chef serves its signature dish to the next ally", "Each Gourmet card serves its line's signature dish (see Dishes sheet) to the next card in lineup order (the previous one if last): Level I dish", "Level II dish", "Level III dish"),
}

KEYWORDS = [
("Initiative bar", "At the start of a duel the bar is at 0 (even); the faster Pokemon moves first (equal speed: lower points, then random). Whenever a Pokemon moves it pushes the bar toward the OTHER side by the opponent's SPEED. The side the bar is on moves next. If the bar lands exactly on 0 it switches to the side that did NOT just move. Equal speeds alternate; speed 2 vs 1 gives the fast one 2 turns per slow turn. The bar carries over when a Pokemon is KO'd; the replacement takes over its side."),
("Turn", "A Pokemon's turn: (1) attack, (2) +1 charge, (3) if charge is full, cast its charge power and reset to 0, (4) resolve its synergy powers. All durations ('2 turns', 'each turn') count the affected Pokemon's OWN turns and tick at the end of its turn."),
("Speed", "PAC speed divided by 7 (cards range about 3-11, most 6-8). A small gap is a small edge: 7 vs 8 gives the faster card about 1 extra turn in 8."),
("Attack", "Deal ATK as PHYSICAL damage. Minimum 1."),
("DEF / SP.DEF", "Flat amount blocked from each PHYSICAL / SPECIAL hit (min 1 damage gets through). TRUE ignores both."),
("SHIELD", "Absorbs damage before HP. Not reduced by DEF. Lost when the card is KO'd."),
("Charge", "Cards cast their charge power when charge reaches Max PP."),
("AP", "Ability power. A number printed with the AP symbol (\u2726) in a charge power grows by +1 for each AP the card has. A number printed \u2726/2 grows by +1 for every 2 AP (rounded down). Plain numbers never change."),
("BURN", "A burning card takes TRUE damage at the end of each of its own turns (1, or 2 from Fire Level III), for 3 turns (4 at Fire Level IV). Burning again resets the timer; the stronger burn wins."),
("POISON", "1 TRUE damage at the end of each of its turns; reapplying adds +1 per turn (stacks) for the rest of that card's time on the field."),
("PARALYZE", "-1 SPEED and its attack gains no charge."),
("FLINCH", "Cannot cast its charge power on its next turn (charge stays full). SILENCE works the same."),
("SLEEP", "Skips its next turn(s) entirely."),
("FREEZE", "Skips its next turn and counts as DEF 0 / SP.DEF 0 until then."),
("CONFUSE", "Its next attack is halved."),
("CHARM", "-2 ATK for 2 turns."),
("WOUND", "Cannot heal or gain SHIELD."),
("ARMOR BREAK", "DEF and SP.DEF count as 0."),
("FATIGUE", "Gains no charge."),
("PROTECT", "Ignores all damage until its next turn begins."),
("SPLASH", "An effect aimed at the OPPONENT'S next entering Pokemon (for example half of an ability's damage, rounded down, hits it when it enters)."),
("BATON", "A bonus given to YOUR OWN next Pokemon, applied when it enters (for example Flora and Sound). The opposite of SPLASH."),
("Crit", "x1.5 damage (round up)."),
("Reveal and order", "At the start of a battle every Pokemon and item on both sides is revealed. Each player sets their lineup order beforehand; Pokemon enter in that order (no choosing replacements). Open question: set before or after seeing the opponent's roster."),
("Synergy level", "Count the cards in the whole lineup with a type (fainted or not), counting each evolution family once per type; level = how many of the type's three thresholds the count reaches (thresholds are the original PAC thresholds x0.75 rounded: see the Synergies sheet; most types have 3 levels, 15 types have 4; Light is set to 2 / 3 / 4). Levels are fixed for the whole battle and apply only to cards that have that type."),
("Scale (provisional)", "HP /10 (min 3), ATK /5 (min 1), DEF and SP.DEF /7, SPEED /7, Max PP /35 (min 2; PAC 50-80 gives 2, 100-120 gives 3), flat powers /10 (percent-of-stat powers: percent x PAC stat /6). Simulated; see duel simulator report."),
]

class C: pass

wb = Workbook()
hdr = PatternFill('solid', fgColor='1F2937'); hf = Font(bold=True, color='FFFFFF')
def sheet(ws, head, rows, widths):
    ws.append(head)
    for c in ws[1]:
        c.fill = hdr; c.font = hf; c.alignment = Alignment(vertical='center')
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions

# cards
import json as _json, os as _os
_TW = _json.load(open(f'{ROOT}/pac-data/tweaks.json')) if _os.path.exists(f'{ROOT}/pac-data/tweaks.json') else {}
rows = []
missing = set()
for c0 in core:
    p = [x for x in allp if x['Index'] == c0['Index'] and x['Name'] == c0['Name']][0]
    stars = int(c0['Stars'])
    c = C()
    c.hp, c.atk, c.defn, c.sdef, c.spd, c.pp = [int(p[k]) for k in ('HP', 'Attack', 'Defense', 'Special Defense', 'Speed', 'Max PP')]
    c.hpc, c.atkc = hp_c(c.hp), atk_c(c.atk)
    ab = p['Ability']
    if ab not in P:
        missing.add(ab); role, txt = '?', ab
    else:
        role, f = P[ab]
        txt = apmark(ab, f(resolve(ab, stars), c))
    types_ = [p[f'Type {i}'] for i in range(1, 5) if p[f'Type {i}']]
    nm = c0['Name'].title().replace('_', ' ')
    tw = _TW.get(c0['Name'].title(), {})
    rows.append([c0['TierGroup'], c0['Index'], nm, stars, float(c0['Value']), int(c0['Cost']),
                 ' / '.join(t.title() for t in types_),
                 max(3, c.hpc + tw.get('hp', 0)), max(1, c.atkc + tw.get('atk', 0)), max(0, df_c(c.defn) + tw.get('df', 0)), max(0, df_c(c.sdef) + tw.get('sdf', 0)),
                 max(1, spd_c(c.spd) + tw.get('spd', 0)), max(1, pp_c(c.pp) + tw.get('pp', 0)),
                 tr['ability'].get(ab, ab), role, txt, ('alt form ' if c0['AltForm'] == 'true' else '') + ('tuned' if tw else '')])
assert not missing, missing
order = {'I': 0, 'II': 1, 'III': 2, 'IV': 3, 'V': 4}
rows.sort(key=lambda r: (order[r[0]], r[4], r[2]))
ws = wb.active; ws.title = 'Cards'
sheet(ws, ['Tier', 'Dex', 'Card', 'Stars', 'Value', 'PAC cost', 'Synergies', 'HP', 'ATK', 'DEF', 'SP.DEF', 'SPEED', 'Max PP',
           'Charge power', 'Role', 'Charge power text (draft)', 'Note'], rows,
      [6, 7, 18, 6, 7, 8, 24, 5, 5, 5, 7, 7, 7, 18, 10, 70, 10])

import sys as _sys; _sys.path.insert(0, f'{ROOT}/pac-sim')
from duel import SYN_TH as _SYN_TH, PAC_TH as _PAC_TH
ws = wb.create_sheet('Synergies')
srows = []
for k, v in SYN.items():
    th = _SYN_TH.get(k.upper(), [2, 3, 4]); pt = _PAC_TH.get(k.upper(), [])
    srows.append([k.title(), ' / '.join(map(str, pt)), ' / '.join(map(str, th)), v[0], v[1], v[2], v[3], v[4] if len(v) > 4 else ''])
sheet(ws, ['Synergy', 'Original PAC thresholds', 'Our thresholds (levels I / II / III / IV)', 'Theme', 'Level I', 'Level II', 'Level III', 'Level IV'], srows, [14, 14, 14, 36, 38, 44, 54, 54])

from duel import DISH as _DISH
DISHN = dict(FARFETCH_D='Leek', GALARIAN_FARFETCH_D='Large Leek', LICKITUNG='Rage Candy Bar', HAPPINY='Nutritious Egg', MUNCHLAX='Leftovers', SHUCKLE='Berry Juice', MILTANK='Moomoo Milk', GULPIN='Black Sludge', SPINDA='Spinda Cocktail', TROPIUS='Nanab Berry', COMBEE='Honey', CHERUBI='Herba Mystica', VANILLITE='Casteliacone', DEERLING_SUMMER='Tea', SINISTEA='Tea', SWIRLIX='Whipped Dream', BOUNSWEET='Fruit Juice', GUZZLORD='Devour', SKWOVET='Berries', APPLIN='Tart Apple', MILCERY='Sweets', LECHONK='Mushrooms', FIDOUGH='Poffin', SMOLIV='Olive Oil', NACLI='Rock Salt', CAPSAKID='Curry', VELUZA='Smoked Filet', DONDOZO='Rice', TATSUGIRI_CURLY='Curly Form', TATSUGIRI_DROOPY='Droopy Form', TATSUGIRI_STRETCHY='Stretchy Form', PECHARUNT='Binding Mochi')
def dish_text(effs):
    out = []
    for e in effs:
        k = e[0]
        if k in ('atk', 'def', 'sdef', 'spd', 'ap'): out.append(f"{e[1]:+d} " + {'atk': 'ATK', 'def': 'DEF', 'sdef': 'SP.DEF', 'spd': 'SPEED', 'ap': 'AP'}[k])
        elif k == 'hp': out.append(f"{e[1]:+d} max HP")
        elif k == 'shield': out.append(f"{e[1]} SHIELD on entry")
        elif k == 'shieldpct': out.append(f"SHIELD worth {e[1]}% max HP (min 2) on entry")
        elif k == 'charge': out.append(f"enters with +{e[1]} charge")
        elif k == 'protect': out.append("PROTECT until its first turn")
        elif k == 'rage': out.append(f"+{e[1]} ATK for its first 3 turns")
        elif k == 'rand': out.append(f"{e[1]} random +1 stat boost(s) (ATK, DEF, SP.DEF, SPEED or +2 max HP)")
        elif k == 'first': out.append(f"its first attack inflicts {e[1].upper()} ({e[2]} turn{'s' if e[2] != 1 else ''})")
        elif k == 'crit': out.append(f"+{e[1]}% crit chance")
        elif k == 'sludge': out.append(f"its attacks POISON the target (up to {e[1]} stack{'s' if e[1] != 1 else ''})")
        elif k == 'nanab': out.append(f"once below 50% HP, heals {e[1]}")
        elif k == 'regen': out.append(f"heals {e[1]} at the end of each of its turns")
        elif k == 'freezen': out.append(f"every {e[1]}th attack FREEZES")
        elif k == 'charmn': out.append(f"every {e[1]}rd attack CHARMS")
        elif k == 'tart': out.append(f"each attack lowers the target's DEF by 1 (max {e[1]} times)")
        elif k == 'dodge': out.append(f"dodges every {e[1]}th attack against it")
    return '; '.join(out)
ws = wb.create_sheet('Gourmet dishes')
drows = []
for fam, (l1, l2, l3) in _DISH.items():
    members = sorted({r['Name'].title().replace('_', ' ') for r in allp if r['Family'] == fam and 'GOURMET' in [r[f'Type {i}'] for i in range(1, 5)]})
    drows.append([fam.title().replace('_', ' '), DISHN.get(fam, ''), ', '.join(members), dish_text(l1), dish_text(l2), dish_text(l3)])
sheet(ws, ['Gourmet line', 'Signature dish', 'Cards', 'Level I dish', 'Level II dish', 'Level III dish'], drows, [22, 18, 50, 50, 50, 50])

ws = wb.create_sheet('Rules & Keywords')
sheet(ws, ['Term', 'Draft rule'], KEYWORDS, [22, 120])

# charge powers by star
ws = wb.create_sheet('Charge powers')
used = collections.defaultdict(list)
for r in rows: used[r[13]].append((r[2], r[3]))
prow = []
for ab, (role, f) in sorted(P.items(), key=lambda kv: tr['ability'].get(kv[0], kv[0])):
    cc = C(); cc.hp = 140; cc.atk = 12; cc.defn = 30; cc.sdef = 30; cc.spd = 50; cc.hpc = 14; cc.atkc = 2
    nm = tr['ability'].get(ab, ab)
    cards = ', '.join(sorted({n for n, s in used.get(nm, [])}))
    prow.append([nm, role, len(used.get(nm, [])), cards])
sheet(ws, ['Power', 'Role', '# cards', 'Used by'], prow, [22, 10, 8, 120])
wb.save(OUT)
print(len(rows), 'cards', len(prow), 'powers')
