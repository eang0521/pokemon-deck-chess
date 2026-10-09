import os as _os
ROOT = _os.environ.get('PAC_ROOT') or _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
import sys, json, os
sys.path.insert(0, f'{ROOT}/pac-sim')
import extra
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
adj = json.load(open(f'{ROOT}/pac-data/tweaks_extra.json')) if os.path.exists(f'{ROOT}/pac-data/tweaks_extra.json') else {}
cs = extra.load_pool(adj, hatch=True)
order = {'I': -1, 'V': 2.5, 'II': 0, 'III': 1, 'IV': 2, 'UNIQUE': 3, 'LEGENDARY': 4}
rows = []
for c in cs:
    rows.append([c.tier, c.dex, c.name, c.stars, round(c.value, 2), '', ' / '.join(t.title() for t in c.torder), c.hp0, c.atk0, c.def0, c.sdef0, c.spd0, c.ppmax,
                 c.abname, c.role, c.text, 'review' if c.flagged else '', c.pool])
rows.sort(key=lambda r: (order[r[0]], r[4], r[2]))
json.dump(rows, open(f'{ROOT}/pac-print/pool_cards.json', 'w'))
wb = Workbook(); ws = wb.active; ws.title = 'Pick cards'
head = ['Tier', 'Dex', 'Card', 'Stars', 'Value', '', 'Synergies', 'HP', 'ATK', 'DEF', 'SP.DEF', 'SPEED', 'Max PP', 'Charge power', 'Role', 'Charge power text', 'Note', 'Pool']
ws.append(head)
for c in ws[1]: c.fill = PatternFill('solid', fgColor='1F2937'); c.font = Font(bold=True, color='FFFFFF')
for r in rows: ws.append(r)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions
for i, w in enumerate([10, 8, 22, 6, 7, 2, 28, 5, 5, 5, 7, 7, 7, 20, 10, 80, 8, 10], 1): ws.column_dimensions[chr(64 + i)].width = w
wb.save(f'{ROOT}/pac-pick-cards-draft.xlsx')
import collections
print(len(rows), collections.Counter((r[17], r[0]) for r in rows))
