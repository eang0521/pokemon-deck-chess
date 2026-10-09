import os as _os
ROOT = _os.environ.get('PAC_ROOT') or _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
import sys, json, os, math, collections
sys.path.insert(0, f'{ROOT}/pac-sim')
import openpyxl
import genpower as G
_orig = G.opsfn; LOG = []
def _ops(parts, c):
    LOG.append(parts); return _orig(parts, c)
G.opsfn = _ops
import engine, extra, duel, items as I
from items import ITEMS, ORDER
_mp = extra.make_power; PARTS = {}
def _mk(*a, **k):
    LOG.clear(); r = _mp(*a, **k); PARTS.setdefault(id(r[2]), []).append(([list(p) for p in LOG[-1]], r[3])); return r
extra.make_power = _mk
adj = json.load(open(f'{ROOT}/pac-data/tweaks_extra.json'))
pool = extra.load_pool(adj, hatch=True)
core = engine.load_cards()
wb = openpyxl.load_workbook(f'{ROOT}/pac-core-cards-draft.xlsx')
txt = {}
for r in wb['Cards'].iter_rows(min_row=2, values_only=True):
    if r[2]: txt[r[2].lower().replace('-', ' ')] = r
def nm(n): return n.lower().replace('-', ' ').replace('_', ' ')
PMIN = {'I': 1, 'II': 3, 'III': 5, 'IV': 8, 'V': 15}; TRADE = {'I': 0, 'II': 2, 'III': 4, 'IV': 7, 'V': 14}
def sy(c): return [t for t in c.torder]
out = []
miss = []
for c in core:
    r = txt.get(nm(c.name))
    if not r: miss.append(c.name)
    out.append(dict(id=len(out), name=c.name, pool='core', tier=c.tier, stars=c.stars, value=c.value, price=math.ceil(c.value - 1e-9), trade=TRADE[c.tier],
        types=sy(c), family=c.family, hp=c.hp0, atk=c.atk0, df=c.def0, sdf=c.sdef0, spd=c.spd0, pp=c.ppmax, ab=c.ability, v=c.v,
        raw=dict(atk=c.raw.atk, spd=c.raw.spd, defn=c.raw.defn, sdef=c.raw.sdef, hpc=c.raw.hpc, atkc=c.raw.atkc), vs=[engine.resolve(c.ability, k) for k in (1,2,3)], abname=(r[13] if r else c.ability), text=((r[15] or '').replace('(scales with SPEED)', '(scales with its current SPEED: every point of SPEED adds damage)') if r else ''), role=(r[14] if r else ''), dex=(r[1] if r else '')))
for c in pool:
    ps = PARTS[id(engine.PW[c.ability])] if id(engine.PW[c.ability]) in PARTS else None
    d = dict(id=len(out), name=c.name, pool=c.pool, tier=c.tier, stars=c.stars, value=c.value, types=sy(c), family=c.family,
        hp=c.hp0, atk=c.atk0, df=c.def0, sdf=c.sdef0, spd=c.spd0, pp=c.ppmax, abname=c.abname, text=c.text, role=c.role, dex=c.dex, ab=c.ability)
    if c.pool == 'add': d['price'] = max(math.ceil(c.value - 1e-9), PMIN.get(c.tier, 1)); d['trade'] = TRADE[c.tier]
    elif c.pool == 'unique': d['price'] = 12; d['trade'] = 12
    elif c.pool == 'legendary': d['price'] = 25; d['trade'] = 25
    else: d['price'] = 0; d['trade'] = TRADE[c.tier]
    out.append(d)

from duel import APM
def parts_of(f):
    if id(f) in PARTS: return PARTS[id(f)][-1]
    return None
for d, c in zip(out[len(core):], pool):
    f = engine.PW[c.ability]
    if c.name == 'Minior':
        cells = [x.cell_contents for x in f.__closure__]
        d['parts2'] = [parts_of(x)[0] for x in cells]
        d['apm2'] = [parts_of(x)[1] for x in cells]
    else:
        p = parts_of(f); d['parts'] = p[0]; 
    d['apm'] = APM.get(c.ability, 'first')
for d, c in zip(out[:len(core)], core): d['apm'] = APM.get(c.ability, 'first')

NAMES = {'Aegislash_Blade':'Aegislash (Blade)','Darmanitan_Zen':'Darmanitan (Zen)','Lycanroc_Day':'Lycanroc (Day)','Lycanroc_Dusk':'Lycanroc (Dusk)','Lycanroc_Night':'Lycanroc (Night)',
 'Meowstic_Female':'Meowstic (F)','Meowstic_Male':'Meowstic (M)','Porygon_2':'Porygon2','Porygon_Z':'Porygon-Z','Ursaluna_Bloodmoon':'Ursaluna (Bloodmoon)','Hakamo O':'Hakamo-o','Jangmo O':'Jangmo-o',
 'Kommo O':'Kommo-o','Ho Oh':'Ho-Oh','Mime Jr':'Mime Jr.','Nidoranf':'Nidoran\u2640','Nidoranm':'Nidoran\u2642','Chi Yu':'Chi-Yu','Oricorio Pa U':"Oricorio (Pa'u)",'Alcremie Vanilla':'Alcremie'}
def clean(n):
    if n in NAMES: return NAMES[n]
    return n.replace('Mr Mime','Mr. Mime').replace('Mr Rime','Mr. Rime').replace('Farfetch D',"Farfetch'd").replace('Sirfetch D',"Sirfetch'd").replace('_',' ')
for d_ in out: d_['name'] = clean(d_['name'])
rows = json.load(open(f'{ROOT}/items_core.json'))['core']
NAME2KEY = {v['name']: k for k, v in ITEMS.items()}
deck = []
for r in rows:
    if r[3] in NAME2KEY: deck.append(dict(key=NAME2KEY[r[3]], tier=r[0], price=r[1], copies=r[2], text=r[5], chips=r[6]))
    else: print('item name missing', r[3])
ORDERL = list(ORDER)
itemsout = {k: dict(key=k, name=v['name'], cat=v['cat'], text=v['text'], stats=v['stats'], flags=v['flags']) for k, v in ITEMS.items()}
syn = {}
for r in wb['Synergies'].iter_rows(min_row=2, values_only=True):
    if r[0]: syn[r[0].upper()] = dict(theme=r[3], lv=[x for x in r[4:8] if x])
dishes = []
for r in wb['Gourmet dishes'].iter_rows(min_row=2, values_only=True):
    if r[0]: dishes.append(dict(line=r[0], dish=r[1], cards=r[2], lv=list(r[3:6])))
DISH = {k: [[list(e) for e in lv] for lv in v] for k, v in duel.DISH.items()}
data = dict(order=ORDERL, cards=out, items=itemsout, itemdeck=deck, syn_th=duel.SYN_TH, syn=syn, dishes=dishes, DISH=DISH, default_dish=[[list(e) for e in lv] for lv in duel.DEFAULT_DISH],
    apm={k: v for k, v in duel.APM.items() if not '#' in k})
os.makedirs(f'{ROOT}/pac-digital/build', exist_ok=True)
js = json.dumps(data, separators=(',', ':'), ensure_ascii=False)
open(f'{ROOT}/pac-digital/build/data.js', 'w').write('window.PAC_DATA=' + js + ';')
json.dump(data, open(f'{ROOT}/pac-digital/build/data.json', 'w'))
print(len(js)//1024, 'KB', len(deck), len(syn), len(dishes))
